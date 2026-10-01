from pathlib import Path

from pypdf import PdfReader
from docx import Document


def extract_text_from_pdf(file_path):
    reader = PdfReader(file_path)

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n".join(pages).strip()


def extract_text_from_docx(file_path):
    document = Document(file_path)

    content = []

    # Read normal paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            content.append(text)

    # Read tables
    for table in document.tables:
        for row in table.rows:
            cells = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    cells.append(cell_text)

            if cells:
                content.append(" | ".join(cells))

    # Read headers
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            text = paragraph.text.strip()

            if text:
                content.append(text)

    # Read footers
    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            text = paragraph.text.strip()

            if text:
                content.append(text)

    return "\n".join(content).strip()


def extract_text_from_txt(file_path):
    path = Path(file_path)

    return path.read_text(
        encoding="utf-8",
        errors="ignore"
    ).strip()


def extract_medical_document_text(file_path):
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"Document not found: {file_path}"
        )

    extension = path.suffix.lower()

    if extension == ".pdf":
        text = extract_text_from_pdf(path)

    elif extension == ".docx":
        text = extract_text_from_docx(path)

    elif extension == ".txt":
        text = extract_text_from_txt(path)

    else:
        raise ValueError(
            "Unsupported document type. "
            "Supported formats: PDF, DOCX, TXT."
        )

    if not text:
        raise ValueError(
            "No readable text was found in the document."
        )

    return text